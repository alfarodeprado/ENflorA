#!/usr/bin/env python3
"""
resolve_accessions.py — fill real ENA accessions into a runs or analysis table.

A run has to reference a registered sample, and an analysis a registered
sample and run. Those accessions only exist once the previous step has been
submitted, so they cannot be typed into the table beforehand.

Instead, write a temporary stand-in in the SAMPLE or RUN_REF column:

    SAMPLE    the 'isolate' you gave the sample in the biosamples table
    RUN_REF   the 'NAME' you gave the read set in the runs table

After submitting the previous step, run this script with the accession file
that step wrote. Values that already look like ENA accessions (SAMEA..., ERS...,
ERR..., and so on) are left alone; everything else is looked up and replaced.
A table may mix both, so objects that were already in ENA can be referenced
directly while new ones are filled in.

Run it from the ENflorA folder, for example:

    python resolve_accessions.py \\
        --table runs/ExperimentList.xlsx \\
        --accessions biosamples/submission/biosample_accessions.txt

    # analysis tables reference both a sample and a run, so give both files
    python resolve_accessions.py \\
        --table analysis/AnalysisList.xlsx \\
        --accessions biosamples/submission/biosample_accessions.txt \\
        --accessions runs/submission/run_accessions.txt

Your table is never modified. An updated copy is written next to it, with
'_resolved' added to the name, and the script tells you which line of
config.yaml to point at it.

Test and live accessions are kept apart: by default the server is taken from
the 'live' key of config.yaml, so a test submission is only ever resolved with
test accessions. Nothing is submitted and ENA is never contacted.
"""

import argparse
import csv
import os
import re
import shutil
import sys

# Accession formats across INSDC:
#   samples   ERS / SRS / DRS, and BioSamples SAMEA / SAMN / SAMD
#   studies   PRJEB / PRJNA / PRJDB, and ERP / SRP / DRP
#   runs      ERR / SRR / DRR      experiments  ERX / SRX / DRX
#   analyses  ERZ / SRZ / DRZ
ACCESSION_RE = re.compile(
    r"^(?:SAM(?:EA|N|D)|PRJ(?:EB|NA|DB)|[ESD]R[SPXRZ])\d+$",
    re.IGNORECASE,
)

# Webin-CLI reports aliases with its context in front: a read set submitted
# with NAME "my_reads" comes back as "webin-reads-my_reads". The user knows the
# NAME, so both forms are accepted.
WEBIN_PREFIX_RE = re.compile(r"^webin-[a-z]+-", re.IGNORECASE)

# Header of each accession file -> (kind, column with the accession, table
# columns it can fill). These headers are written by biosamples.py, runs.py
# and analysis.py.
ACCESSION_FILES = {
    "accession\talias\tserver": ("sample", "accession", ["SAMPLE"]),
    "experiment_accession\trun_accession\talias\tserver": ("run", "run_accession", ["RUN_REF"]),
    "analysis_accession\talias\tserver": ("analysis", "analysis_accession", []),
}

# Only RUN_REF may hold several references (an assembly built from several
# read sets). A SAMPLE value is never split, so a lab code containing a comma
# stays intact.
LIST_COLUMNS = {"RUN_REF"}
LIST_SEPARATOR = ","

TABLE_EXTENSIONS = {".xlsx", ".csv", ".tsv", ".tab", ".txt"}


def looks_like_accession(value: str) -> bool:
    return bool(ACCESSION_RE.match(value.strip()))


def cell_text(value) -> str:
    """Text of a cell as the user sees it; 12.0 from a spreadsheet becomes '12'."""
    if value is None:
        return ""
    if isinstance(value, float) and value.is_integer():
        return str(int(value))
    return str(value).strip()


def default_server(config_path: str) -> str:
    """'live' if config.yaml says live: True, otherwise 'test'."""
    if not os.path.exists(config_path):
        return "test"
    try:
        import yaml
        with open(config_path) as handle:
            config = yaml.safe_load(handle) or {}
    except Exception as exc:
        print(f"WARNING: could not read {config_path} ({exc}); assuming the test server.")
        return "test"
    return "live" if config.get("live") is True else "test"


def read_accession_file(path: str, server: str):
    """
    Read an accession file written by one of the submission scripts.

    Returns (kind, target_columns, lookup, other_server), where lookup maps a
    lowercased temporary accession to the real one for the chosen server, and
    other_server holds the rows from the other server, so a near miss can be
    reported clearly instead of just 'not found'.
    """
    if not os.path.exists(path):
        sys.exit(f"Accession file not found: {path}")

    with open(path, encoding="utf-8") as handle:
        lines = [line.rstrip("\r\n") for line in handle if line.strip()]
    if not lines:
        sys.exit(f"Accession file is empty: {path}")

    header = lines[0]
    if header not in ACCESSION_FILES:
        if header.endswith("\talias"):
            sys.exit(
                f"'{path}' was written by an older version of ENflorA and has no "
                "'server' column, so test and live accessions cannot be told "
                "apart.\nRun the submission script once more with the current "
                "version (it upgrades the file), or add a last column reading "
                "'test' or 'live' to every line by hand."
            )
        sys.exit(
            f"'{path}' is not an ENflorA accession file.\nExpected a first line of:\n  "
            + "\n  ".join(h.replace("\t", "<tab>") for h in ACCESSION_FILES)
        )

    kind, accession_column, target_columns = ACCESSION_FILES[header]
    fields = header.split("\t")
    accession_index = fields.index(accession_column)
    alias_index = fields.index("alias")
    server_index = fields.index("server")

    lookup, other_server = {}, {}
    for number, line in enumerate(lines[1:], start=2):
        parts = line.split("\t")
        if len(parts) != len(fields):
            print(f"WARNING: skipping malformed line {number} of {path}")
            continue
        accession = parts[accession_index].strip()
        alias = parts[alias_index].strip()
        row_server = parts[server_index].strip().lower()
        if not accession or accession.lower() == "none":
            continue
        target = lookup if row_server == server else other_server

        for key in {alias, WEBIN_PREFIX_RE.sub("", alias)}:
            folded = key.lower()
            if folded in target and target[folded] != accession:
                sys.exit(
                    f"'{path}' gives two different accessions for '{key}' on the "
                    f"{row_server} server ({target[folded]} and {accession}). "
                    "Fix the file before resolving."
                )
            target[folded] = accession

    # Each accession is stored under two spellings of its alias, so count
    # accessions rather than keys.
    print(f"Read {len(set(lookup.values()))} {kind} accession(s) from the "
          f"{server} server in {path}")
    if other_server:
        print(f"  ({len(set(other_server.values()))} from the other server are ignored)")
    return kind, target_columns, lookup, other_server


def resolve_cell(value, column, lookup, other_server, server, stats, location):
    """
    Resolve one cell. Returns the new text, or None to leave it exactly as it was.
    """
    text = cell_text(value)
    if not text:
        return None

    if column in LIST_COLUMNS:
        parts = [part.strip() for part in text.split(LIST_SEPARATOR) if part.strip()]
    else:
        parts = [text]

    resolved, changed = [], False
    for part in parts:
        if looks_like_accession(part):
            stats["kept"] += 1
            resolved.append(part)
            continue

        match = lookup.get(part.lower())
        if match:
            stats["substituted"] += 1
            resolved.append(match)
            changed = True
            continue

        if part.lower() in other_server:
            other = "live" if server == "test" else "test"
            stats["problems"].append(
                f"{location}: '{part}' was registered on the {other} server "
                f"(as {other_server[part.lower()]}), not on {server}"
            )
        else:
            stats["problems"].append(
                f"{location}: '{part}' is neither an ENA accession nor in the accession file(s)"
            )
        resolved.append(part)

    if not changed:
        return None
    return (LIST_SEPARATOR + " ").join(resolved)


def find_columns(header, wanted):
    """Map each wanted column to its index in the header, ignoring case and spaces."""
    normalised = [cell_text(h).upper() for h in header]
    return {name: normalised.index(name) for name in wanted if name in normalised}


def missing_columns_error(table, wanted, header):
    listed = ", ".join(cell_text(h) for h in header if cell_text(h))
    sys.exit(
        f"None of the column(s) to fill ({', '.join(wanted)}) are in '{table}'.\n"
        f"Its columns are: {listed}"
    )


def resolve_xlsx(in_path, out_path, lookups, server, stats):
    """Edit the first worksheet of a copy, keeping every other sheet intact."""
    from openpyxl import load_workbook

    shutil.copyfile(in_path, out_path)
    workbook = load_workbook(out_path)
    sheet = workbook.worksheets[0]

    header = [cell.value for cell in sheet[1]]
    indexes = find_columns(header, lookups)
    if not indexes:
        os.remove(out_path)
        missing_columns_error(in_path, list(lookups), header)

    for name, index in indexes.items():
        lookup, other_server = lookups[name]
        for row in range(2, sheet.max_row + 1):
            cell = sheet.cell(row=row, column=index + 1)
            new = resolve_cell(cell.value, name, lookup, other_server, server,
                               stats, f"{name}, row {row - 1}")
            if new is not None:
                cell.value = new

    workbook.save(out_path)
    return header, list(indexes)


def resolve_delimited(in_path, out_path, lookups, server, stats, separator):
    with open(in_path, newline="", encoding="utf-8-sig") as handle:
        rows = list(csv.reader(handle, delimiter=separator))
    if not rows:
        sys.exit(f"'{in_path}' is empty.")

    header = rows[0]
    indexes = find_columns(header, lookups)
    if not indexes:
        missing_columns_error(in_path, list(lookups), header)

    for number, row in enumerate(rows[1:], start=1):
        for name, index in indexes.items():
            if index >= len(row):
                continue
            lookup, other_server = lookups[name]
            new = resolve_cell(row[index], name, lookup, other_server, server,
                               stats, f"{name}, row {number}")
            if new is not None:
                row[index] = new

    with open(out_path, "w", newline="", encoding="utf-8") as handle:
        csv.writer(handle, delimiter=separator).writerows(rows)
    return header, list(indexes)


def next_step_hint(header, output):
    """Tell the user which config.yaml key should point at the new table."""
    names = {cell_text(h).upper() for h in header}
    if "RUN_REF" in names or "ASSEMBLYNAME" in names:
        key, folder = "data_analysis", "analysis"
    else:
        key, folder = "data_runs", "runs"
    here = os.path.dirname(os.path.abspath(__file__))
    relative = os.path.relpath(os.path.abspath(output), os.path.join(here, folder))
    return (
        f"Next: in config.yaml, set\n"
        f"    {key}: {relative}\n"
        f"then run {folder}.py as usual."
    )


def main():
    here = os.path.dirname(os.path.abspath(__file__))
    parser = argparse.ArgumentParser(
        description="Replace temporary stand-ins (sample isolate, run NAME) in a "
                    "runs or analysis table with the real ENA accessions assigned "
                    "by the previous submission step.")
    parser.add_argument("--table", required=True,
                        help="Runs or analysis table to fill in (.xlsx, .csv, .tsv/.tab/.txt)")
    parser.add_argument("--accessions", required=True, action="append", metavar="FILE",
                        help="Accession file from biosamples.py or runs.py "
                             "(biosample_accessions.txt, run_accessions.txt). "
                             "Give it twice to use both.")
    parser.add_argument("-o", "--output",
                        help="Where to write the updated table "
                             "(default: next to the input, with '_resolved' added)")
    parser.add_argument("--server", choices=["test", "live"],
                        help="Which server's accessions to use. Default: taken "
                             "from 'live' in config.yaml (test unless live: True).")
    parser.add_argument("--config", default=os.path.join(here, "config.yaml"),
                        help="Config file to read 'live' from (default: config.yaml "
                             "in the ENflorA folder)")
    parser.add_argument("--column", action="append", metavar="NAME",
                        help="Column to fill, if not the usual one (SAMPLE for "
                             "sample accessions, RUN_REF for run accessions). "
                             "Only with a single --accessions file.")
    parser.add_argument("--allow-missing", action="store_true",
                        help="Write the table even if some values cannot be "
                             "resolved, leaving those values as they are")
    parser.add_argument("--force", action="store_true",
                        help="Overwrite the output file if it already exists")
    args = parser.parse_args()

    # --- The table -------------------------------------------------------
    if not os.path.exists(args.table):
        sys.exit(f"Table not found: {args.table}")
    extension = os.path.splitext(args.table)[1].lower()
    if extension == ".xls":
        sys.exit("Old-style .xls files are not supported here. Save the table "
                 "as .xlsx (or .csv) and run again.")
    if extension not in TABLE_EXTENSIONS:
        sys.exit(f"Unsupported table extension '{extension}'. "
                 "Use .xlsx, .csv, or .tsv/.tab/.txt")

    output = args.output
    if not output:
        stem, ext = os.path.splitext(args.table)
        output = f"{stem}_resolved{ext}"
    if os.path.abspath(output) == os.path.abspath(args.table):
        sys.exit("The output would overwrite your original table. Choose "
                 "another name with -o.")
    if os.path.splitext(output)[1].lower() != extension:
        sys.exit(f"The output must have the same extension as the input ({extension}).")
    if os.path.exists(output) and not args.force:
        sys.exit(f"'{output}' already exists. Use -o for another name, or --force "
                 "to overwrite it.")

    # --- The accession files -------------------------------------------
    if args.server:
        server = args.server
        print(f"Using {server} server accessions (--server).")
    else:
        server = default_server(args.config)
        print(f"Using {server} server accessions (from 'live' in "
              f"{os.path.relpath(args.config)}; change with --server).")

    if args.column and len(args.accessions) > 1:
        sys.exit("--column can only be used with a single --accessions file.")

    lookups = {}
    for path in args.accessions:
        kind, targets, lookup, other_server = read_accession_file(path, server)
        targets = [c.strip().upper() for c in args.column] if args.column else targets
        if not targets:
            sys.exit(
                f"Nothing in a runs or analysis table refers to {kind} accessions, "
                f"so '{path}' has no column to fill. Use --column to name one."
            )
        for column in targets:
            if column in lookups:
                sys.exit(f"Two accession files both fill {column}; give only one.")
            lookups[column] = (lookup, other_server)

    # --- Resolve ---------------------------------------------------------
    stats = {"substituted": 0, "kept": 0, "problems": []}
    if extension == ".xlsx":
        header, used = resolve_xlsx(args.table, output, lookups, server, stats)
    else:
        separator = "," if extension == ".csv" else "\t"
        header, used = resolve_delimited(args.table, output, lookups, server,
                                         stats, separator)

    print(f"\nColumn(s) checked: {', '.join(used)}")
    unused = [c for c in lookups if c not in used]
    if unused:
        print(f"  (not in this table, skipped: {', '.join(unused)})")
    print(f"  {stats['substituted']} temporary value(s) replaced with accessions")
    print(f"  {stats['kept']} value(s) already ENA accessions, left as they were")

    problems = stats["problems"]
    if problems:
        print(f"  {len(problems)} value(s) could not be resolved:")
        for problem in problems:
            print(f"    {problem}")

    if problems and not args.allow_missing:
        if os.path.exists(output):
            os.remove(output)
        sys.exit(
            "\nNothing was written. Usually this means the previous step has not "
            "been submitted yet, or a stand-in is spelled differently from the "
            "isolate / NAME it was submitted under.\nTo write the table anyway, "
            "leaving those values untouched, add --allow-missing."
        )

    print(f"\nUpdated table written to: {output}")
    if problems:
        print("Some values were left unresolved. Fix them before submitting.")
    print(next_step_hint(header, output))


if __name__ == "__main__":
    main()
